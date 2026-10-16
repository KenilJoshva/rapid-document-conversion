"""Unit tests for the document conversion engine."""

import io
import pytest
from PIL import Image
import docx
import pypdf

from conversion_app.converter import (
    convert_document,
    ConversionError,
    _text_to_pdf,
    _markdown_to_docx,
    _docx_to_markdown,
)


def test_txt_to_pdf_conversion():
    sample_text = "Rapid Document Conversion\nAutomated quality validation."
    pdf_bytes, mime = convert_document(sample_text.encode("utf-8"), "txt", "pdf")

    assert mime == "application/pdf"
    assert pdf_bytes.startswith(b"%PDF")
    reader = pypdf.PdfReader(io.BytesIO(pdf_bytes))
    assert len(reader.pages) >= 1
    extracted = reader.pages[0].extract_text()
    assert "Rapid Document Conversion" in extracted


def test_txt_to_docx_conversion():
    sample_text = "Paragraph 1\n\nParagraph 2"
    docx_bytes, mime = convert_document(sample_text.encode("utf-8"), "txt", "docx")

    assert "wordprocessingml" in mime
    assert docx_bytes.startswith(b"PK\x03\x04")
    doc = docx.Document(io.BytesIO(docx_bytes))
    assert len(doc.paragraphs) >= 2


def test_txt_to_html_and_md():
    text = "Line 1\nLine 2"
    html_bytes, html_mime = convert_document(text.encode("utf-8"), "txt", "html")
    assert html_mime == "text/html"
    assert b"Line 1" in html_bytes

    md_bytes, md_mime = convert_document(text.encode("utf-8"), "txt", "md")
    assert md_mime == "text/markdown"
    assert md_bytes.decode("utf-8") == text


def test_markdown_to_docx_and_html():
    md = "# Main Header\n\n- Item 1\n- Item 2\n\nRegular text here."
    docx_bytes, _ = convert_document(md.encode("utf-8"), "md", "docx")
    assert docx_bytes.startswith(b"PK\x03\x04")
    doc = docx.Document(io.BytesIO(docx_bytes))
    assert any("Main Header" in p.text for p in doc.paragraphs)

    html_bytes, _ = convert_document(md.encode("utf-8"), "md", "html")
    assert b"<h1>Main Header</h1>" in html_bytes
    assert b"<li>Item 1</li>" in html_bytes


def test_docx_to_markdown_and_txt():
    # Build a test docx in memory
    doc = docx.Document()
    doc.add_heading("Doc Heading", level=1)
    doc.add_paragraph("This is content in docx.")
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    md_bytes, _ = convert_document(docx_bytes, "docx", "md")
    md_str = md_bytes.decode("utf-8")
    assert "# Doc Heading" in md_str
    assert "This is content in docx." in md_str

    txt_bytes, _ = convert_document(docx_bytes, "docx", "txt")
    txt_str = txt_bytes.decode("utf-8")
    assert "Doc Heading" in txt_str
    assert "This is content in docx." in txt_str


def test_pdf_to_txt_and_md():
    sample_text = "Confidential Report Page One\nDetailed Analysis."
    pdf_bytes = _text_to_pdf(sample_text)

    txt_bytes, _ = convert_document(pdf_bytes, "pdf", "txt")
    extracted_text = txt_bytes.decode("utf-8")
    assert "Confidential Report Page One" in extracted_text

    md_bytes, _ = convert_document(pdf_bytes, "pdf", "md")
    assert b"# Extracted PDF Content" in md_bytes
    assert b"Confidential Report Page One" in md_bytes


def test_image_transcoding():
    # Create test 100x100 PNG
    img = Image.new("RGBA", (100, 100), color=(255, 0, 0, 255))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    png_bytes = buf.getvalue()

    # Convert to JPEG (requires alpha stripping handled internally)
    jpg_bytes, mime = convert_document(png_bytes, "png", "jpg")
    assert mime == "image/jpeg"
    assert jpg_bytes.startswith(b"\xff\xd8\xff")

    # Convert to WebP
    webp_bytes, webp_mime = convert_document(png_bytes, "png", "webp")
    assert webp_mime == "image/webp"
    assert webp_bytes.startswith(b"RIFF")
    assert b"WEBP" in webp_bytes[:12]

    # Convert Image to PDF
    img_pdf, pdf_mime = convert_document(png_bytes, "png", "pdf")
    assert pdf_mime == "application/pdf"
    assert img_pdf.startswith(b"%PDF")


def test_unsupported_conversion_raises():
    with pytest.raises(ConversionError) as exc_info:
        convert_document(b"random-data", "pdf", "docx")
    assert "unsupported" in str(exc_info.value).lower()
