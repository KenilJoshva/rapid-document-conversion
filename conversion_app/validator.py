"""Automatic quality validation module.

Validates conversion success, magic file signatures, file sizes, readability,
and extracts/compares content between source and target documents.
"""

import io
from typing import Dict, Any, List, Optional
from PIL import Image
import pypdf
import docx

# Magic byte signatures for verified binary formats
FILE_SIGNATURES = {
    "pdf": [b"%PDF"],
    "docx": [b"PK\x03\x04"],
    "png": [b"\x89PNG\r\n\x1a\n"],
    "jpg": [b"\xff\xd8\xff"],
    "jpeg": [b"\xff\xd8\xff"],
    "webp": [b"RIFF"],  # also checks WEBP chunk at offset 8
    "bmp": [b"BM"],
}


def verify_file_signature(data: bytes, expected_format: str) -> bool:
    """Verify that file data matches expected format magic bytes."""
    fmt = expected_format.lower().strip()
    if fmt in ("txt", "md", "html"):
        # Text formats: check valid UTF-8 and mostly printable characters
        try:
            text = data.decode("utf-8")
            # If empty, fails
            if not text:
                return False
            printable_count = sum(1 for c in text if c.isprintable() or c in "\r\n\t")
            ratio = printable_count / len(text)
            return ratio >= 0.85
        except UnicodeDecodeError:
            return False

    signatures = FILE_SIGNATURES.get(fmt, [])
    if not signatures:
        # Unknown binary format, fallback to non-empty
        return len(data) > 0

    if fmt == "webp":
        return len(data) >= 12 and data.startswith(b"RIFF") and data[8:12] == b"WEBP"

    return any(data.startswith(sig) for sig in signatures)


def extract_text_sample(data: bytes, fmt: str, max_chars: int = 5000) -> str:
    """Extract plain text from document bytes for validation comparison."""
    fmt = fmt.lower().strip()
    try:
        if fmt in ("txt", "md", "html"):
            return data.decode("utf-8", errors="ignore")[:max_chars]

        if fmt == "docx":
            doc = docx.Document(io.BytesIO(data))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            return "\n".join(paragraphs)[:max_chars]

        if fmt == "pdf":
            reader = pypdf.PdfReader(io.BytesIO(data))
            text_parts = []
            for page in reader.pages[:10]:
                text = page.extract_text()
                if text:
                    text_parts.append(text)
            return "\n".join(text_parts)[:max_chars]
    except Exception:
        return ""
    return ""


def validate_conversion_quality(
    input_data: bytes,
    input_format: str,
    output_data: bytes,
    output_format: str,
) -> Dict[str, Any]:
    """Execute quality validation suite on conversion input and output.

    Returns
    -------
    dict
        Structured validation report with status, passed/warn/fail checks,
        and diagnostic metrics.
    """
    checks: List[Dict[str, str]] = []
    overall_status = "PASS"
    metrics: Dict[str, Any] = {
        "input_bytes": len(input_data),
        "output_bytes": len(output_data),
        "input_format": input_format.lower(),
        "output_format": output_format.lower(),
    }

    # 1. Output existence and non-zero size
    if not output_data or len(output_data) == 0:
        checks.append({
            "name": "File Size Non-Zero",
            "status": "FAIL",
            "message": "Output document is empty (0 bytes).",
        })
        return {
            "status": "FAIL",
            "checks": checks,
            "metrics": metrics,
            "summary": "Conversion failed: produced empty output.",
        }

    checks.append({
        "name": "File Size Non-Zero",
        "status": "PASS",
        "message": f"Generated {len(output_data):,} bytes successfully.",
    })

    # 2. Magic signature check
    sig_valid = verify_file_signature(output_data, output_format)
    if sig_valid:
        checks.append({
            "name": "File Signature / Magic Bytes",
            "status": "PASS",
            "message": f"Verified valid file header and structure for {output_format.upper()}.",
        })
    else:
        checks.append({
            "name": "File Signature / Magic Bytes",
            "status": "FAIL",
            "message": f"Output file does not match expected signature for {output_format.upper()}.",
        })
        overall_status = "FAIL"

    # 3. Readability and syntax parsing validation
    out_fmt = output_format.lower()
    readability_passed = True
    readability_msg = ""

    try:
        if out_fmt == "pdf":
            reader = pypdf.PdfReader(io.BytesIO(output_data))
            page_count = len(reader.pages)
            metrics["page_count"] = page_count
            if page_count > 0:
                readability_msg = f"PDF is fully readable with {page_count} page(s)."
            else:
                readability_passed = False
                readability_msg = "PDF reader could not identify any pages."
        elif out_fmt == "docx":
            doc = docx.Document(io.BytesIO(output_data))
            p_count = len(doc.paragraphs)
            metrics["paragraph_count"] = p_count
            readability_msg = f"DOCX document successfully parsed with {p_count} paragraph(s)."
        elif out_fmt in ("png", "jpg", "jpeg", "webp", "bmp"):
            with Image.open(io.BytesIO(output_data)) as img:
                metrics["dimensions"] = f"{img.width}x{img.height}"
                metrics["mode"] = img.mode
                readability_msg = f"Image readable: {img.width}x{img.height} ({img.mode})."
        elif out_fmt in ("txt", "md", "html"):
            decoded = output_data.decode("utf-8")
            metrics["character_count"] = len(decoded)
            metrics["line_count"] = decoded.count("\n") + 1
            readability_msg = f"Text document decoded cleanly with {len(decoded):,} character(s)."
    except Exception as e:
        readability_passed = False
        readability_msg = f"Readability check failed: {str(e)}"

    if readability_passed:
        checks.append({
            "name": "Format Readability",
            "status": "PASS",
            "message": readability_msg,
        })
    else:
        checks.append({
            "name": "Format Readability",
            "status": "FAIL",
            "message": readability_msg,
        })
        overall_status = "FAIL"

    # 4. Text comparison (when both input and output are text-bearing)
    in_fmt = input_format.lower()
    text_bearing_formats = ("txt", "md", "docx", "pdf", "html")
    if in_fmt in text_bearing_formats and out_fmt in text_bearing_formats:
        in_text = extract_text_sample(input_data, in_fmt)
        out_text = extract_text_sample(output_data, out_fmt)

        metrics["input_extracted_chars"] = len(in_text)
        metrics["output_extracted_chars"] = len(out_text)

        if len(in_text.strip()) > 20 and len(out_text.strip()) == 0:
            checks.append({
                "name": "Text Preservation",
                "status": "WARNING",
                "message": "Input contained text, but no extractable text was found in output.",
            })
            if overall_status == "PASS":
                overall_status = "WARNING"
        elif len(in_text.strip()) > 0 and len(out_text.strip()) > 0:
            # Word token overlap check
            in_words = set(in_text.lower().split()[:200])
            out_words = set(out_text.lower().split()[:200])
            common = in_words.intersection(out_words)
            overlap_ratio = len(common) / max(len(in_words), 1)
            metrics["text_token_overlap"] = round(overlap_ratio, 2)

            if overlap_ratio >= 0.3:
                checks.append({
                    "name": "Text Preservation",
                    "status": "PASS",
                    "message": f"Content preserved accurately (token overlap: {round(overlap_ratio * 100)}%).",
                })
            else:
                checks.append({
                    "name": "Text Preservation",
                    "status": "WARNING",
                    "message": f"Significant textual divergence observed (overlap: {round(overlap_ratio * 100)}%).",
                })
                if overall_status == "PASS":
                    overall_status = "WARNING"

    # 5. Size sanity check
    ratio = len(output_data) / max(len(input_data), 1)
    metrics["size_ratio"] = round(ratio, 3)

    if ratio < 0.01 and len(input_data) > 1024:
        checks.append({
            "name": "Size Sanity Check",
            "status": "WARNING",
            "message": f"Output file is very small relative to input ({len(output_data)} bytes vs {len(input_data)} bytes).",
        })
        if overall_status == "PASS":
            overall_status = "WARNING"
    elif ratio > 50.0 and len(output_data) > 5 * 1024 * 1024:
        checks.append({
            "name": "Size Sanity Check",
            "status": "WARNING",
            "message": f"Output file experienced significant expansion ({round(ratio, 1)}x).",
        })
        if overall_status == "PASS":
            overall_status = "WARNING"
    else:
        checks.append({
            "name": "Size Sanity Check",
            "status": "PASS",
            "message": f"Output size is normal ({round(ratio, 2)}x relative to source).",
        })

    # Collect explicit issues explaining detected warnings or failures
    issues: List[Dict[str, str]] = []
    for c in checks:
        if c["status"] in ("WARNING", "FAIL"):
            rec_text = "Review source document formatting or select an alternative output format."
            if c["name"] == "File Size Non-Zero":
                rec_text = "Ensure the source document contains valid renderable content."
            elif c["name"] == "File Signature / Magic Bytes":
                rec_text = f"The generated bytes did not form a valid {out_fmt.upper()} header structure. Try re-converting or selecting a different target format."
            elif c["name"] == "Format Readability":
                rec_text = f"The {out_fmt.upper()} reader could not parse the output. Check if source document contains unsupported encoding or macros."
            elif c["name"] == "Text Preservation":
                rec_text = "If source document contains images or scanned pages, standard text extraction cannot capture OCR text. Ensure source uses digital text layers."
            elif c["name"] == "Size Sanity Check":
                rec_text = "Check whether graphics or uncompressed elements caused abnormal file sizing."

            issues.append({
                "check": c["name"],
                "severity": c["status"],
                "issue": c["message"],
                "recommendation": rec_text,
            })

    summary_text = (
        "All validation checks passed with high fidelity."
        if overall_status == "PASS"
        else f"Validation completed with {len(issues)} notice(s)/warning(s)."
        if overall_status == "WARNING"
        else f"Quality validation detected critical errors ({len(issues)} failure(s))."
    )

    return {
        "status": overall_status,
        "checks": checks,
        "issues": issues,
        "metrics": metrics,
        "summary": summary_text,
    }

